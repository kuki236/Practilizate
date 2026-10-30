import csv
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.data_source import AxDataSource, StrRef
from openpyxl.chart.marker import DataPoint
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from practilizate.logging import get_logger
from practilizate.models import Offer, QualityReport

LOGGER = get_logger("practilizate.export")

STATUS_ES = {
    "active": "Vigente",
    "expiring_soon": "Por vencer",
    "expired": "Vencida",
    "unknown": "Sin fecha",
}
STATUS_ORDER = ("Vigente", "Por vencer", "Vencida", "Sin fecha")
SOURCE_ES = {"fixtures": "Fixtures sintéticos", "live": "practicas.pe (en vivo)"}
RULE_DESCRIPTIONS_ES = {
    "R01": "Campos obligatorios: título, empresa y URL",
    "R02": "URL de oferta única",
    "R03": "URL absoluta válida",
    "R04": "Fecha de publicación válida",
    "R05": "Fecha de vencimiento válida",
    "R06": "Vencimiento no anterior a la publicación",
    "R07": "Subvención numérica o nula explícita",
    "R08": "Subvención dentro de los límites configurados",
    "R09": "Empresa no vacía después de normalizar",
    "R10": "Ubicación presente",
    "R11": "Modalidad en el conjunto permitido",
    "R12": "Asistencia en el conjunto permitido",
    "R13": "Al menos una carrera asociada",
    "R14": "Título, empresa y vencimiento duplicados",
    "R15": "Habilidades presentes con requisitos detallados",
    "R16": "Marca de tiempo de la fuente presente",
    "R17": "Texto válido en UTF-8 y normalizado",
}

NAVY = "1F4E79"
BLUE = "2E75B6"
LIGHT_BLUE = "DDEBF7"
INK = "1F2937"
MUTED = "6B7280"
GREEN = "548235"
LIGHT_GREEN = "E2EFDA"
AMBER = "BF8F00"
LIGHT_AMBER = "FFF2CC"
RED = "C00000"
LIGHT_RED = "FCE4E4"
WHITE = "FFFFFF"

THIN = Side(style="thin", color="D1D5DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

CSV_FIELDS = [
    "offer_id",
    "url",
    "title",
    "company",
    "company_normalized",
    "location",
    "district",
    "department",
    "modality",
    "attendance",
    "salary_text",
    "salary_amount",
    "salary_currency",
    "published_at",
    "deadline_at",
    "status",
    "career_count",
    "careers",
    "skill_count",
    "skills",
    "requirements_text",
    "source_fetched_at",
    "quality_score",
    "days_to_deadline",
]


@dataclass
class ExportContext:
    run_id: str
    source: str
    as_of: date
    generated_at: datetime
    runtime_seconds: float
    quality: QualityReport
    salary_max: float = 12000.0


@dataclass
class ExportResult:
    csv_rows: int
    excel_sheets: int


def _ordered_offers(offers: list[Offer]) -> list[Offer]:
    return sorted(
        offers,
        key=lambda offer: (
            (0, int(offer.offer_id)) if offer.offer_id.isdigit() else (1, offer.offer_id)
        ),
    )


def _careers(offer: Offer) -> str:
    return "; ".join(career.name for career in offer.careers)


def _skills(offer: Offer) -> str:
    return "; ".join(offer.skills)


def _days_to_deadline(offer: Offer, as_of: date) -> int | None:
    if offer.deadline_at is None:
        return None
    return (offer.deadline_at - as_of).days


def write_dashboard_csv(offers: list[Offer], path: Path, as_of: date) -> int:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=CSV_FIELDS, lineterminator="\n")
        writer.writeheader()
        for offer in _ordered_offers(offers):
            writer.writerow(
                {
                    "offer_id": offer.offer_id,
                    "url": offer.url,
                    "title": offer.title,
                    "company": offer.company or "",
                    "company_normalized": offer.company_normalized or "",
                    "location": offer.location or "",
                    "district": offer.district or "",
                    "department": offer.department or "",
                    "modality": offer.modality or "",
                    "attendance": offer.attendance or "",
                    "salary_text": offer.salary_text or "",
                    "salary_amount": offer.salary_amount if offer.salary_amount is not None else "",
                    "salary_currency": offer.salary_currency or "",
                    "published_at": offer.published_at.isoformat() if offer.published_at else "",
                    "deadline_at": offer.deadline_at.isoformat() if offer.deadline_at else "",
                    "status": offer.status,
                    "career_count": len(offer.careers),
                    "careers": _careers(offer),
                    "skill_count": len(offer.skills),
                    "skills": _skills(offer),
                    "requirements_text": offer.requirements_text,
                    "source_fetched_at": offer.source_fetched_at.isoformat(timespec="seconds"),
                    "quality_score": offer.quality_score,
                    "days_to_deadline": _days_to_deadline(offer, as_of),
                }
            )
    LOGGER.info("dashboard csv written", extra={"path": str(path), "rows": len(offers)})
    return len(offers)


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def _header_style(cell) -> None:
    cell.font = Font(name="Calibri", size=11, bold=True, color=WHITE)
    cell.fill = _fill(NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BORDER


def _title(ws, row: int, title: str, subtitle: str | None = None, end_col: int = 10) -> None:
    cell = ws.cell(row=row, column=1, value=title)
    cell.font = Font(name="Calibri", size=20, bold=True, color=NAVY)
    cell.alignment = Alignment(vertical="center")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=end_col)
    ws.row_dimensions[row].height = 30
    if subtitle:
        sub = ws.cell(row=row + 1, column=1, value=subtitle)
        sub.font = Font(name="Calibri", size=11, color=MUTED)
        sub.alignment = Alignment(vertical="center", wrap_text=True)
        ws.merge_cells(start_row=row + 1, start_column=1, end_row=row + 1, end_column=end_col)


def _add_table(
    ws,
    name: str,
    headers: list[str],
    rows: list[list[Any]],
    *,
    start_row: int,
    start_col: int = 1,
    style: str = "TableStyleMedium2",
    number_formats: list[str | None] | None = None,
) -> tuple[int, int, int, int]:
    for offset, header in enumerate(headers):
        cell = ws.cell(row=start_row, column=start_col + offset, value=header)
        _header_style(cell)
    for row_offset, values in enumerate(rows, start=1):
        for col_offset, value in enumerate(values):
            cell = ws.cell(row=start_row + row_offset, column=start_col + col_offset, value=value)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=False)
            if number_formats and number_formats[col_offset]:
                cell.number_format = number_formats[col_offset]
    end_row = start_row + len(rows)
    end_col = start_col + len(headers) - 1
    ref = f"{get_column_letter(start_col)}{start_row}:{get_column_letter(end_col)}{end_row}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name=style, showFirstColumn=False, showLastColumn=False, showRowStripes=True
    )
    ws.add_table(table)
    return start_row, start_col, end_row, end_col


def _widths(ws, widths: dict[str, int]) -> None:
    for column, width in widths.items():
        ws.column_dimensions[column].width = width


def _use_text_categories(chart, ws, labels_ref: Reference) -> None:
    column = get_column_letter(labels_ref.min_col)
    formula = f"'{ws.title}'!${column}${labels_ref.min_row}:${column}${labels_ref.max_row}"
    chart.series[0].cat = AxDataSource(strRef=StrRef(f=formula))


def _bar_chart(
    ws,
    title: str,
    labels_ref: Reference,
    values_ref: Reference,
    *,
    anchor: str,
    horizontal: bool = False,
    color: str = BLUE,
    width: float = 14.0,
    height: float = 8.0,
) -> BarChart:
    chart = BarChart()
    chart.type = "bar" if horizontal else "col"
    chart.style = 10
    chart.title = title
    chart.width = width
    chart.height = height
    chart.add_data(values_ref, titles_from_data=True)
    chart.set_categories(labels_ref)
    _use_text_categories(chart, ws, labels_ref)
    chart.legend = None
    chart.y_axis.title = None
    chart.x_axis.title = None
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.x_axis.tickLblPos = "nextTo"
    chart.y_axis.tickLblPos = "nextTo"
    chart.y_axis.majorGridlines = None
    if horizontal:
        chart.x_axis.axPos = "l"
        chart.y_axis.axPos = "b"
    else:
        chart.x_axis.axPos = "b"
        chart.y_axis.axPos = "l"
    series = chart.series[0]
    series.graphicalProperties.solidFill = color
    series.graphicalProperties.line.solidFill = color
    ws.add_chart(chart, anchor)
    return chart


def _pie_chart(
    ws, title: str, labels_ref: Reference, values_ref: Reference, *, anchor: str
) -> PieChart:
    chart = PieChart()
    chart.title = title
    chart.width = 12.0
    chart.height = 8.0
    chart.add_data(values_ref, titles_from_data=True)
    chart.set_categories(labels_ref)
    _use_text_categories(chart, ws, labels_ref)
    chart.legend.position = "r"
    colors = [GREEN, AMBER, RED, MUTED]
    for index, color in enumerate(colors):
        point = DataPoint(idx=index)
        point.graphicalProperties.solidFill = color
        chart.series[0].data_points.append(point)
    ws.add_chart(chart, anchor)
    return chart


def _career_counts(offers: list[Offer]) -> Counter[str]:
    counts: Counter[str] = Counter()
    for offer in offers:
        for career in offer.careers:
            counts[career.name] += 1
    return counts


def _company_rows(offers: list[Offer], salary_max: float) -> list[list[Any]]:
    grouped: dict[str, list[Offer]] = defaultdict(list)
    for offer in offers:
        grouped[offer.company or "Sin empresa"].append(offer)
    rows = []
    for company, items in grouped.items():
        salaries = [
            item.salary_amount
            for item in items
            if item.salary_amount is not None and item.salary_amount <= salary_max
        ]
        rows.append(
            [
                company,
                len(items),
                sum(item.quality_score for item in items) / len(items),
                sum(salaries) / len(salaries) if salaries else None,
            ]
        )
    return sorted(rows, key=lambda row: (-row[1], str(row[0])))


def _skill_rows(offers: list[Offer]) -> list[list[Any]]:
    counts = Counter(skill for offer in offers for skill in offer.skills)
    ordered = sorted(counts.items(), key=lambda item: (-item[1], item[0]))
    return [
        [skill, count, f"=B{5 + index} / COUNTA(tblDatos[ID oferta])"]
        for index, (skill, count) in enumerate(ordered)
    ]


def _salary_rows(
    offers: list[Offer], career_counts: Counter[str], salary_max: float
) -> list[list[Any]]:
    rows = []
    for career, _ in career_counts.most_common(10):
        items = [offer for offer in offers if career in {c.name for c in offer.careers}]
        salaries = [
            offer.salary_amount
            for offer in items
            if offer.salary_amount is not None and offer.salary_amount <= salary_max
        ]
        rows.append(
            [
                career,
                len(salaries),
                min(salaries) if salaries else None,
                sum(salaries) / len(salaries) if salaries else None,
                max(salaries) if salaries else None,
            ]
        )
    return rows


def _status_rows(offers: list[Offer]) -> list[list[Any]]:
    counts = Counter(STATUS_ES.get(offer.status, "Sin fecha") for offer in offers)
    return [[status, counts.get(status, 0)] for status in STATUS_ORDER]


def _quality_rows(quality: QualityReport) -> list[list[Any]]:
    severity_es = {"critical": "Crítica", "warning": "Advertencia"}
    return [
        [
            rule.rule_id,
            RULE_DESCRIPTIONS_ES.get(rule.rule_id, rule.description),
            severity_es.get(rule.severity, rule.severity),
            rule.checked,
            rule.failed,
            rule.pass_rate,
            ", ".join(rule.sample_failed_offer_ids),
        ]
        for rule in quality.rules
    ]


def _dictionary_rows() -> list[list[str]]:
    return [
        ["ID oferta", "Texto", "Identificador público de la convocatoria en practicas.pe.", "URL"],
        ["Título", "Texto", "Título publicado de la oferta.", "Listado y detalle"],
        ["Empresa", "Texto", "Entidad que publica la convocatoria.", "Listado y detalle"],
        ["Ubicación", "Texto", "Distrito y departamento donde se realiza la práctica.", "Detalle"],
        ["Distrito", "Texto", "Distrito reportado por la fuente.", "Detalle"],
        ["Departamento", "Texto", "Departamento reportado por la fuente.", "Detalle"],
        ["Modalidad", "Texto", "Preprofesional o profesional.", "Filtro y detalle"],
        ["Asistencia", "Texto", "Presencial, híbrido o virtual.", "Detalle"],
        ["Subvención", "Texto", "Texto original de la subvención publicada.", "Detalle"],
        ["Monto", "Número", "Subvención mensual en soles; vacío si no se publica.", "Detalle"],
        ["Publicado", "Fecha", "Fecha de publicación de la oferta.", "Detalle"],
        ["Vence", "Fecha", "Fecha límite para postular.", "Detalle"],
        ["Días para vencer", "Entero", "Días entre la fecha de corte y el vencimiento.", "Cálculo"],
        ["Estado", "Texto", "Vigente, por vencer, vencida o sin fecha.", "Cálculo"],
        [
            "Carreras",
            "Texto",
            "Carreras asociadas, separadas por punto y coma.",
            "Filtro de búsqueda",
        ],
        ["Habilidades", "Texto", "Habilidades detectadas en los requisitos.", "Normalización"],
        [
            "Calidad",
            "Porcentaje",
            "Proporción de reglas de calidad aprobadas por la oferta.",
            "Motor de calidad",
        ],
        ["URL", "Texto", "Enlace público a la convocatoria.", "Fuente"],
        ["Requisitos", "Texto", "Requisitos y condiciones, sin datos personales.", "Detalle"],
    ]


def write_powerbi_support_csvs(
    offers: list[Offer], context: ExportContext, output_dir: Path
) -> dict[str, int]:
    output_dir.mkdir(parents=True, exist_ok=True)
    ordered = _ordered_offers(offers)
    counts: dict[str, int] = {}

    career_rows = [
        {
            "offer_id": offer.offer_id,
            "career_id": career.id,
            "career_name": career.name,
        }
        for offer in ordered
        for career in offer.careers
    ]
    career_path = output_dir / "offer_careers.csv"
    with career_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(
            handle, fieldnames=["offer_id", "career_id", "career_name"], lineterminator="\n"
        )
        writer.writeheader()
        writer.writerows(career_rows)
    counts["offer_careers"] = len(career_rows)

    skill_rows = [
        {"offer_id": offer.offer_id, "skill": skill} for offer in ordered for skill in offer.skills
    ]
    skill_path = output_dir / "offer_skills.csv"
    with skill_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["offer_id", "skill"], lineterminator="\n")
        writer.writeheader()
        writer.writerows(skill_rows)
    counts["offer_skills"] = len(skill_rows)

    quality_rows = [
        {
            "rule_id": rule.rule_id,
            "description": rule.description,
            "description_es": RULE_DESCRIPTIONS_ES.get(rule.rule_id, rule.description),
            "severity": rule.severity,
            "severity_es": "Crítica" if rule.severity == "critical" else "Advertencia",
            "checked": rule.checked,
            "failed": rule.failed,
            "pass_rate": rule.pass_rate,
            "sample_failed_offer_ids": "; ".join(rule.sample_failed_offer_ids),
        }
        for rule in context.quality.rules
    ]
    quality_path = output_dir / "quality_rules.csv"
    with quality_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=[
                "rule_id",
                "description",
                "description_es",
                "severity",
                "severity_es",
                "checked",
                "failed",
                "pass_rate",
                "sample_failed_offer_ids",
            ],
            lineterminator="\n",
        )
        writer.writeheader()
        writer.writerows(quality_rows)
    counts["quality_rules"] = len(quality_rows)

    salaries = [offer.salary_amount for offer in ordered if offer.salary_amount is not None]
    metadata = {
        "run_id": context.run_id,
        "source": context.source,
        "as_of": context.as_of.isoformat(),
        "generated_at": context.generated_at.isoformat(timespec="seconds"),
        "runtime_seconds": context.runtime_seconds,
        "offers": len(ordered),
        "companies": len({offer.company for offer in ordered if offer.company}),
        "careers": len({career.id for offer in ordered for career in offer.careers}),
        "skills": len({skill for offer in ordered for skill in offer.skills}),
        "quality_score": context.quality.overall_score,
        "expiring": sum(1 for offer in ordered if offer.status == "expiring_soon"),
        "salary_average": round(sum(salaries) / len(salaries), 2) if salaries else 0,
        "last_updated": context.generated_at.strftime("%d/%m/%Y %H:%M"),
    }
    metadata_path = output_dir / "run_metadata.csv"
    with metadata_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(metadata), lineterminator="\n")
        writer.writeheader()
        writer.writerow(metadata)
    counts["run_metadata"] = 1

    language_path = output_dir / "dim_idioma.csv"
    with language_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(
            handle, fieldnames=["Idioma", "Etiqueta", "Orden"], lineterminator="\n"
        )
        writer.writeheader()
        writer.writerows(
            [
                {"Idioma": "ES", "Etiqueta": "Español", "Orden": 1},
                {"Idioma": "EN", "Etiqueta": "English", "Orden": 2},
            ]
        )
    counts["dim_idioma"] = 2

    status_path = output_dir / "dim_estado.csv"
    with status_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=["status", "label_es", "label_en", "orden"],
            lineterminator="\n",
        )
        writer.writeheader()
        writer.writerows(
            [
                {"status": "active", "label_es": "Vigente", "label_en": "Active", "orden": 1},
                {
                    "status": "expiring_soon",
                    "label_es": "Por vencer",
                    "label_en": "Expiring soon",
                    "orden": 2,
                },
                {"status": "expired", "label_es": "Vencida", "label_en": "Expired", "orden": 3},
                {"status": "unknown", "label_es": "Sin fecha", "label_en": "No date", "orden": 4},
            ]
        )
    counts["dim_estado"] = 4

    LOGGER.info("power bi support csvs written", extra={"path": str(output_dir), "counts": counts})
    return counts


def write_excel(offers: list[Offer], context: ExportContext, path: Path) -> int:
    path.parent.mkdir(parents=True, exist_ok=True)
    ordered = _ordered_offers(offers)
    career_counts = _career_counts(ordered)
    company_counts = Counter(offer.company or "Sin empresa" for offer in ordered)
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = "Practilizate"
    wb.properties.title = "Practilizate — Observatorio de prácticas preprofesionales"
    wb.properties.subject = "Ofertas de prácticas preprofesionales en Lima"
    wb.properties.description = (
        "Datos públicos de practicas.pe. Practilizate no está afiliado a practicas.pe."
    )
    wb.properties.keywords = "prácticas, Perú, empleabilidad, datos abiertos"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True

    careers_by_count = [name for name, _ in career_counts.most_common(10)]
    skills_by_count = [row[0] for row in _skill_rows(ordered)]

    listas = wb.create_sheet("Listas")
    listas.sheet_state = "hidden"
    listas["A1"] = "Carreras"
    listas["B1"] = "Habilidades"
    listas["C1"] = "Empresas"
    listas["D1"] = "Estados"
    for row_index, career in enumerate(careers_by_count, start=2):
        listas.cell(row=row_index, column=1, value=career)
    for row_index, skill in enumerate(skills_by_count, start=2):
        listas.cell(row=row_index, column=2, value=skill)
    for row_index, company in enumerate(
        [name for name in company_counts if name != "Sin empresa"], start=2
    ):
        listas.cell(row=row_index, column=3, value=company)
    for row_index, status in enumerate(STATUS_ORDER, start=2):
        listas.cell(row=row_index, column=4, value=status)
    _widths(listas, {"A": 34, "B": 30, "C": 34, "D": 16})

    portada = wb.create_sheet("Portada")
    _build_portada(portada, context=context)

    datos = wb.create_sheet("Datos")
    _build_datos(datos, ordered, context.as_of)

    resumen = wb.create_sheet("Resumen")
    _build_resumen(resumen, ordered, career_counts=career_counts, salary_max=context.salary_max)

    habilidades = wb.create_sheet("Habilidades")
    _build_habilidades(
        habilidades,
        ordered,
        skill_rows=_skill_rows(ordered),
        careers_by_count=careers_by_count,
    )

    empresas = wb.create_sheet("Empresas")
    _build_empresas(
        empresas,
        company_rows=_company_rows(ordered, context.salary_max),
        salary_rows=_salary_rows(ordered, career_counts, context.salary_max),
    )

    calidad = wb.create_sheet("Calidad")
    _build_calidad(calidad, context.quality)

    diccionario = wb.create_sheet("Diccionario")
    _build_diccionario(diccionario)

    wb.defined_names["ListaCarreras"] = DefinedName(
        "ListaCarreras",
        attr_text=f"Listas!$A$2:$A${max(2, len(careers_by_count) + 1)}",
    )

    wb.move_sheet("Portada", offset=-wb.sheetnames.index("Portada"))
    wb.move_sheet("Listas", offset=len(wb.sheetnames) - 1 - wb.sheetnames.index("Listas"))
    wb.active = 0
    temp_path = path.with_suffix(".xlsx.tmp")
    wb.save(temp_path)
    temp_path.replace(path)
    LOGGER.info(
        "excel workbook written",
        extra={"path": str(path), "sheets": len(wb.sheetnames), "offers": len(ordered)},
    )
    return len(wb.sheetnames)


def _build_portada(ws, *, context: ExportContext) -> None:
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = NAVY
    _title(
        ws,
        1,
        "Practilizate",
        "Observatorio de prácticas preprofesionales del Perú — datos públicos de practicas.pe",
        end_col=8,
    )
    ws["A4"] = "Fuente"
    ws["B4"] = SOURCE_ES.get(context.source, context.source)
    ws["A5"] = "Fecha de corte"
    ws["B5"] = context.as_of
    ws["B5"].number_format = "dd/mm/yyyy"
    ws["A6"] = "Última actualización"
    ws["B6"] = context.generated_at
    ws["B6"].number_format = "dd/mm/yyyy hh:mm"
    ws["A7"] = "Identificador del run"
    ws["B7"] = context.run_id
    ws["A8"] = "Tiempo de pipeline"
    ws["B8"] = context.runtime_seconds
    ws["B8"].number_format = '0.00 "s"'

    kpis = [
        ("Ofertas", "='Resumen'!B6", "0"),
        ("Empresas", "='Resumen'!C6", "0"),
        ("Carreras", "='Resumen'!D6", "0"),
        ("Habilidades", "='Resumen'!E6", "0"),
        ("Calidad promedio", "='Resumen'!F6", "0.0%"),
        ("Vencen en 7 días", "='Resumen'!G6", "0"),
        ("Subvención promedio", "='Resumen'!H6", '"S/ "#,##0.00'),
    ]
    for index, (label, formula, number_format) in enumerate(kpis, start=1):
        col = index
        label_cell = ws.cell(row=10, column=col, value=label)
        label_cell.font = Font(name="Calibri", size=10, bold=True, color=MUTED)
        label_cell.alignment = Alignment(wrap_text=True, vertical="bottom")
        value_cell = ws.cell(row=11, column=col, value=formula)
        value_cell.font = Font(name="Calibri", size=18, bold=True, color=NAVY)
        value_cell.number_format = number_format
        value_cell.fill = _fill(LIGHT_BLUE)
        value_cell.border = BORDER
        value_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[11].height = 28
    ws.conditional_formatting.add(
        "E11",
        ColorScaleRule(
            start_type="num",
            start_value=0,
            start_color=RED,
            end_type="num",
            end_value=1,
            end_color=GREEN,
        ),
    )

    ws["A14"] = "Cómo leer este libro"
    ws["A14"].font = Font(name="Calibri", size=12, bold=True, color=NAVY)
    guide = [
        ("Datos", "Una fila por oferta, con carreras y habilidades separadas por punto y coma."),
        ("Resumen", "KPIs, filtro por carrera y gráficos de estado."),
        ("Habilidades", "Habilidades más pedidas y su demanda por carrera."),
        ("Empresas", "Empresas con más ofertas y rango de subvenciones por carrera."),
        ("Calidad", "Resultado de las 17 reglas; las reglas críticas detienen el pipeline."),
        ("Diccionario", "Definición de cada campo y su origen."),
    ]
    for offset, (sheet, description) in enumerate(guide, start=15):
        ws.cell(row=offset, column=1, value=sheet).font = Font(bold=True, color=INK)
        ws.cell(row=offset, column=2, value=description).font = Font(color=INK)
        ws.merge_cells(start_row=offset, start_column=2, end_row=offset, end_column=8)
        ws.row_dimensions[offset].height = 20

    ws["A23"] = "Atribución"
    ws["A23"].font = Font(bold=True, color=NAVY)
    ws["A24"] = (
        "Datos públicos de practicas.pe. Practilizate no está afiliado a practicas.pe. "
        "Solo se publican datos agregados; no se incluyen datos personales."
    )
    ws["A24"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A24:H25")
    ws["A26"] = (
        "Limitaciones: cobertura de Lima y modalidad preprofesional (v0.1); la subvención "
        "puede no estar publicada; los datos se reportan sin ocultar errores de calidad."
    )
    ws["A26"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A26:H27")
    _widths(ws, {"A": 22, "B": 18, "C": 16, "D": 16, "E": 18, "F": 18, "G": 20, "H": 18})


def _build_datos(ws, offers: list[Offer], as_of: date) -> None:
    ws.sheet_properties.tabColor = BLUE
    _title(
        ws, 1, "Datos", "Una fila por oferta normalizada. Filtra con los encabezados.", end_col=19
    )
    headers = [
        "ID oferta",
        "Título",
        "Empresa",
        "Ubicación",
        "Distrito",
        "Departamento",
        "Modalidad",
        "Asistencia",
        "Subvención",
        "Monto",
        "Publicado",
        "Vence",
        "Días para vencer",
        "Estado",
        "Carreras",
        "Habilidades",
        "Calidad",
        "URL",
        "Requisitos",
    ]
    rows: list[list[Any]] = []
    for offer in offers:
        rows.append(
            [
                int(offer.offer_id) if offer.offer_id.isdigit() else offer.offer_id,
                offer.title,
                offer.company or "No publicada",
                offer.location or "No publicada",
                offer.district or "",
                offer.department or "",
                offer.modality or "",
                offer.attendance or "",
                offer.salary_text or "No publicada",
                offer.salary_amount,
                offer.published_at,
                offer.deadline_at,
                _days_to_deadline(offer, as_of),
                STATUS_ES.get(offer.status, "Sin fecha"),
                _careers(offer),
                _skills(offer),
                offer.quality_score,
                offer.url,
                offer.requirements_text,
            ]
        )
    _start, _, end, _ = _add_table(
        ws,
        "tblDatos",
        headers,
        rows,
        start_row=4,
        number_formats=[
            "0",
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            '"S/ "#,##0.00',
            "dd/mm/yyyy",
            "dd/mm/yyyy",
            "0",
            None,
            None,
            None,
            "0.0%",
            None,
            None,
        ],
    )
    ws.freeze_panes = ws.cell(row=5, column=2)
    _widths(
        ws,
        {
            "A": 10,
            "B": 42,
            "C": 22,
            "D": 22,
            "E": 16,
            "F": 14,
            "G": 14,
            "H": 12,
            "I": 16,
            "J": 12,
            "K": 12,
            "L": 12,
            "M": 10,
            "N": 13,
            "O": 32,
            "P": 36,
            "Q": 10,
            "R": 60,
            "S": 70,
        },
    )
    ws.conditional_formatting.add(
        f"N5:N{end}",
        FormulaRule(formula=['$N5="Vencida"'], fill=_fill(LIGHT_RED), font=Font(color=RED)),
    )
    ws.conditional_formatting.add(
        f"N5:N{end}",
        FormulaRule(formula=['$N5="Por vencer"'], fill=_fill(LIGHT_AMBER), font=Font(color=AMBER)),
    )
    ws.conditional_formatting.add(
        f"N5:N{end}",
        FormulaRule(formula=['$N5="Vigente"'], fill=_fill(LIGHT_GREEN), font=Font(color=GREEN)),
    )
    ws.conditional_formatting.add(
        f"Q5:Q{end}",
        ColorScaleRule(
            start_type="num",
            start_value=0,
            start_color=RED,
            end_type="num",
            end_value=1,
            end_color=GREEN,
        ),
    )


def _build_resumen(
    ws,
    offers: list[Offer],
    *,
    career_counts: Counter[str],
    salary_max: float,
) -> None:
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = NAVY
    _title(
        ws,
        1,
        "Resumen ejecutivo",
        "Indicadores del mercado de prácticas. Usa el filtro para explorar una carrera.",
        end_col=10,
    )
    ws["A3"] = "Filtrar por carrera:"
    ws["A3"].font = Font(bold=True, color=INK)
    first_career = [name for name, _ in career_counts.most_common()]
    ws["C3"] = first_career[0] if first_career else ""
    ws["C3"].fill = _fill(LIGHT_BLUE)
    ws["C3"].border = BORDER
    ws["C3"].font = Font(bold=True, color=NAVY)
    dv = DataValidation(type="list", formula1="=ListaCarreras", allow_blank=False)
    dv.error = "Selecciona una carrera de la lista."
    dv.errorTitle = "Carrera no válida"
    dv.prompt = "Elige una carrera para filtrar el indicador de la fila 9."
    dv.promptTitle = "Filtro de carrera"
    ws.add_data_validation(dv)
    dv.add("C3")

    labels = [
        "Ofertas",
        "Empresas",
        "Carreras",
        "Habilidades",
        "Calidad promedio",
        "Vencen en 7 días",
        "Subvención promedio",
    ]
    formulas = [
        "=COUNTA(tblDatos[ID oferta])",
        "=COUNTA(Listas!$C$2:$C$200)",
        "=COUNTA(Listas!$A$2:$A$200)",
        "=COUNTA(tblHabilidades[Habilidad])",
        "='Calidad'!B3",
        '=COUNTIF(tblDatos[Estado],"Por vencer")',
        f'=IFERROR(AVERAGEIFS(tblDatos[Monto],tblDatos[Monto],">0",tblDatos[Monto],"<={int(salary_max)}"),0)',
    ]
    formats = ["0", "0", "0", "0", "0.0%", "0", '"S/ "#,##0.00']
    for index, (label, formula, number_format) in enumerate(
        zip(labels, formulas, formats, strict=True), start=2
    ):
        label_cell = ws.cell(row=5, column=index, value=label)
        label_cell.font = Font(size=10, bold=True, color=MUTED)
        label_cell.alignment = Alignment(wrap_text=True, vertical="bottom")
        value_cell = ws.cell(row=6, column=index, value=formula)
        value_cell.font = Font(size=18, bold=True, color=NAVY)
        value_cell.fill = _fill(LIGHT_BLUE)
        value_cell.border = BORDER
        value_cell.number_format = number_format
        value_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[5].height = 30
    ws.row_dimensions[6].height = 30
    ws["B8"] = "Ofertas de la carrera seleccionada"
    ws["B8"].font = Font(size=10, bold=True, color=MUTED)
    ws["C9"] = '=COUNTIF(tblDatos[Carreras],"*"&$C$3&"*")'
    ws["C9"].font = Font(size=18, bold=True, color=BLUE)
    ws["C9"].number_format = "0"
    ws["H7"] = f"Promedios sin valores atípicos > S/ {salary_max:,.0f}"
    ws["H7"].font = Font(size=9, color=MUTED)
    ws["D9"] = "Las ofertas con más de una carrera se cuentan en cada una."
    ws["D9"].font = Font(size=10, color=MUTED)
    ws.merge_cells("D9:J9")

    career_rows = []
    for index, (career, _) in enumerate(career_counts.most_common(10)):
        row_ref = f"$A{31 + index}"
        career_rows.append(
            [
                career,
                f'=COUNTIF(tblDatos[Carreras],"*"&{row_ref}&"*")',
                f'=IFERROR(AVERAGEIF(tblDatos[Carreras],"*"&{row_ref}&"*",tblDatos[Calidad]),"")',
                f'=IFERROR(AVERAGEIFS(tblDatos[Monto],tblDatos[Carreras],"*"&{row_ref}&"*",tblDatos[Monto],">0",tblDatos[Monto],"<={int(salary_max)}"),"")',
            ]
        )
    start, _, end, _ = _add_table(
        ws,
        "tblPorCarrera",
        ["Carrera", "Ofertas", "Calidad promedio", "Subvención promedio"],
        career_rows,
        start_row=30,
        number_formats=[None, "0", "0.0%", '"S/ "#,##0.00'],
    )
    status_rows = _status_rows(offers)
    _add_table(
        ws,
        "tblPorEstado",
        ["Estado", "Ofertas"],
        status_rows,
        start_row=30,
        start_col=7,
        number_formats=[None, "0"],
    )

    labels_ref = Reference(ws, min_col=1, min_row=start + 1, max_row=end)
    values_ref = Reference(ws, min_col=2, min_row=start, max_row=end)
    _bar_chart(
        ws,
        "Ofertas por carrera",
        labels_ref,
        values_ref,
        anchor="B12",
        horizontal=True,
        color=BLUE,
        width=15.0,
        height=7.5,
    )
    status_end = 30 + len(status_rows)
    _pie_chart(
        ws,
        "Estado de las ofertas",
        Reference(ws, min_col=7, min_row=31, max_row=status_end),
        Reference(ws, min_col=8, min_row=30, max_row=status_end),
        anchor="H12",
    )
    _widths(
        ws,
        {
            "A": 34,
            "B": 14,
            "C": 22,
            "D": 24,
            "E": 16,
            "F": 16,
            "G": 18,
            "H": 18,
            "I": 14,
            "J": 14,
        },
    )
    ws.conditional_formatting.add(
        "F6",
        ColorScaleRule(
            start_type="num",
            start_value=0,
            start_color=RED,
            end_type="num",
            end_value=1,
            end_color=GREEN,
        ),
    )


def _build_habilidades(
    ws,
    offers: list[Offer],
    *,
    skill_rows: list[list[Any]],
    careers_by_count: list[str],
) -> None:
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = GREEN
    _title(
        ws,
        1,
        "Habilidades",
        "Demanda de habilidades en las ofertas. Se cuentan ofertas, no menciones.",
        end_col=12,
    )
    start, _, end, _ = _add_table(
        ws,
        "tblHabilidades",
        ["Habilidad", "Ofertas", "% de ofertas"],
        skill_rows,
        start_row=4,
        number_formats=[None, "0", "0.0%"],
    )
    labels_ref = Reference(ws, min_col=1, min_row=start + 1, max_row=min(end, start + 10))
    values_ref = Reference(ws, min_col=2, min_row=start, max_row=min(end, start + 10))
    _bar_chart(
        ws,
        "Top 10 habilidades",
        labels_ref,
        values_ref,
        anchor="E4",
        horizontal=True,
        color=GREEN,
        width=16.0,
        height=10.0,
    )

    top_careers = careers_by_count[:6]
    matrix_rows: list[list[Any]] = []
    for skill in [row[0] for row in skill_rows[:10]]:
        row: list[Any] = [skill]
        for career in top_careers:
            row.append(
                sum(
                    1
                    for offer in offers
                    if skill in offer.skills and career in {item.name for item in offer.careers}
                )
            )
        row.append(sum(row[1:]))
        matrix_rows.append(row)
    matrix_start = 30
    header = ["Habilidad", *top_careers, "Total"]
    for offset, value in enumerate(header):
        cell = ws.cell(row=matrix_start, column=1 + offset, value=value)
        _header_style(cell)
    for row_offset, row in enumerate(matrix_rows, start=1):
        ws.cell(row=matrix_start + row_offset, column=1, value=row[0]).border = BORDER
        for col_offset in range(2, len(header)):
            cell = ws.cell(
                row=matrix_start + row_offset,
                column=col_offset,
                value=row[col_offset - 1],
            )
            cell.border = BORDER
            cell.number_format = "0"
    last_col = get_column_letter(len(header))
    table = Table(
        displayName="tblHabilidadesCarrera",
        ref=f"A{matrix_start}:{last_col}{matrix_start + len(matrix_rows)}",
    )
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=True)
    ws.add_table(table)
    ws.cell(row=matrix_start - 1, column=1, value="Top habilidades por carrera").font = Font(
        size=12, bold=True, color=NAVY
    )
    note = ws.cell(
        row=matrix_start + len(matrix_rows) + 2,
        column=1,
        value="Cada celda cuenta ofertas que mencionan la habilidad en esa carrera. Una oferta puede aparecer en varias carreras.",
    )
    note.font = Font(size=10, color=MUTED)
    ws.merge_cells(
        start_row=matrix_start + len(matrix_rows) + 2,
        start_column=1,
        end_row=matrix_start + len(matrix_rows) + 2,
        end_column=len(header),
    )
    _widths(ws, {"A": 34, "B": 18, "C": 18, "D": 18, "E": 18, "F": 18, "G": 18, "H": 12})


def _build_empresas(
    ws,
    *,
    company_rows: list[list[Any]],
    salary_rows: list[list[Any]],
) -> None:
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = AMBER
    _title(
        ws,
        1,
        "Empresas y subvenciones",
        "Concentración de ofertas y rangos de subvención por carrera.",
        end_col=12,
    )
    start, _, end, _ = _add_table(
        ws,
        "tblEmpresas",
        ["Empresa", "Ofertas", "Calidad promedio", "Subvención promedio"],
        company_rows[:15],
        start_row=4,
        number_formats=[None, "0", "0.0%", '"S/ "#,##0.00'],
    )
    labels_ref = Reference(ws, min_col=1, min_row=start + 1, max_row=min(end, start + 10))
    values_ref = Reference(ws, min_col=2, min_row=start, max_row=min(end, start + 10))
    _bar_chart(
        ws,
        "Top 10 empresas por ofertas",
        labels_ref,
        values_ref,
        anchor="G4",
        horizontal=True,
        color=AMBER,
        width=16.0,
        height=10.0,
    )

    salary_start = 30
    salary_table_start, _, salary_end, _ = _add_table(
        ws,
        "tblSalariosCarrera",
        ["Carrera", "Ofertas con subvención", "Mínima", "Promedio", "Máxima"],
        salary_rows,
        start_row=salary_start,
        number_formats=[None, "0", '"S/ "#,##0.00', '"S/ "#,##0.00', '"S/ "#,##0.00'],
    )
    _bar_chart(
        ws,
        "Subvención promedio por carrera",
        Reference(ws, min_col=1, min_row=salary_table_start + 1, max_row=salary_end),
        Reference(ws, min_col=4, min_row=salary_table_start, max_row=salary_end),
        anchor="G30",
        horizontal=True,
        color=NAVY,
        width=16.0,
        height=10.0,
    )
    _widths(ws, {"A": 38, "B": 16, "C": 18, "D": 20, "E": 16, "F": 14, "G": 18, "H": 14})


def _build_calidad(ws, quality: QualityReport) -> None:
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = RED
    _title(
        ws,
        1,
        "Calidad de datos",
        "17 reglas. Las reglas críticas detienen la publicación si fallan.",
        end_col=12,
    )
    ws["A3"] = "Calidad general"
    ws["A3"].font = Font(bold=True, color=MUTED)
    ws["A3"].alignment = Alignment(wrap_text=True, vertical="center")
    ws["B3"] = "=AVERAGE(tblCalidad[Tasa de aprobación])"
    ws["B3"].number_format = "0.0%"
    ws["B3"].font = Font(size=18, bold=True, color=NAVY)
    ws["C3"] = "Reglas fallidas"
    ws["C3"].font = Font(bold=True, color=MUTED)
    ws["C3"].alignment = Alignment(wrap_text=True, vertical="center")
    ws["D3"] = '=COUNTIF(tblCalidad[Fallidas],">0")'
    ws["D3"].font = Font(
        size=18, bold=True, color=RED if any(rule.failed for rule in quality.rules) else GREEN
    )
    ws["D3"].number_format = "0"
    ws["E3"] = "Reglas críticas"
    ws["E3"].font = Font(bold=True, color=MUTED)
    ws["E3"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[3].height = 30
    ws["F3"] = ", ".join(quality.critical_failures) if quality.critical_failures else "Ninguna"
    ws["F3"].font = Font(bold=True, color=RED if quality.critical_failures else GREEN)

    start, _, end, _ = _add_table(
        ws,
        "tblCalidad",
        [
            "Regla",
            "Descripción",
            "Severidad",
            "Revisadas",
            "Fallidas",
            "Tasa de aprobación",
            "Ejemplos",
        ],
        _quality_rows(quality),
        start_row=5,
        number_formats=[None, None, None, "0", "0", "0.0%", None],
    )
    ws.conditional_formatting.add(
        f"E6:E{end}",
        CellIsRule(
            operator="greaterThan", formula=["0"], fill=_fill(LIGHT_RED), font=Font(color=RED)
        ),
    )
    ws.conditional_formatting.add(
        f"F6:F{end}",
        ColorScaleRule(
            start_type="num",
            start_value=0,
            start_color=RED,
            end_type="num",
            end_value=1,
            end_color=GREEN,
        ),
    )
    _bar_chart(
        ws,
        "Ofertas fallidas por regla",
        Reference(ws, min_col=1, min_row=start + 1, max_row=end),
        Reference(ws, min_col=5, min_row=start, max_row=end),
        anchor="I5",
        horizontal=True,
        color=RED,
        width=16.0,
        height=12.0,
    )
    ws.freeze_panes = "A6"
    _widths(ws, {"A": 12, "B": 56, "C": 16, "D": 12, "E": 14, "F": 22, "G": 22})


def _build_diccionario(ws) -> None:
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = MUTED
    _title(
        ws,
        1,
        "Diccionario de datos",
        "Definiciones de los campos del libro y de la fuente.",
        end_col=6,
    )
    _add_table(
        ws,
        "tblDiccionario",
        ["Campo", "Tipo", "Descripción", "Origen"],
        _dictionary_rows(),
        start_row=4,
        number_formats=[None, None, None, None],
    )
    ws.freeze_panes = "A5"
    _widths(ws, {"A": 26, "B": 14, "C": 76, "D": 24})
