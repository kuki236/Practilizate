import csv
import json
import zipfile
from pathlib import Path

import openpyxl

CSV_FIELDS = [
    "offer_id",
    "url",
    "title",
    "company",
    "company_normalized",
    "location",
    "district",
    "department",
    "modality",
    "attendance",
    "salary_text",
    "salary_amount",
    "salary_currency",
    "published_at",
    "deadline_at",
    "status",
    "career_count",
    "careers",
    "skill_count",
    "skills",
    "requirements_text",
    "source_fetched_at",
    "quality_score",
    "days_to_deadline",
]


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def test_dashboard_csv_matches_schema_and_contract(e2e_run: Path) -> None:
    path = e2e_run / "dashboard_dataset.csv"
    assert path.is_file()
    assert path.read_bytes().startswith(b"\xef\xbb\xbf")
    rows = read_csv(path)
    assert len(rows) == 12
    assert list(rows[0].keys()) == CSV_FIELDS

    by_id = {row["offer_id"]: row for row in rows}
    assert by_id["90001"]["career_count"] == "2"
    assert by_id["90001"]["careers"] == "ADMINISTRACION; INGENIERIA DE SISTEMAS"
    assert by_id["90001"]["skill_count"] == "5"
    assert by_id["90002"]["salary_amount"] == ""
    assert by_id["90002"]["days_to_deadline"] == "4"
    assert by_id["90006"]["company"] == ""
    assert by_id["90008"]["status"] == "expired"
    assert int(by_id["90008"]["days_to_deadline"]) < 0
    assert by_id["90009"]["location"] == ""
    assert by_id["90011"]["skills"] == ""
    assert "reclutamiento@example.com" not in by_id["90003"]["requirements_text"]
    assert "999 111 222" not in by_id["90003"]["requirements_text"]


def test_excel_workbook_is_complete_and_well_formed(e2e_run: Path) -> None:
    path = e2e_run / "practilizate.xlsx"
    assert path.is_file()
    with zipfile.ZipFile(path) as archive:
        assert archive.testzip() is None

    workbook = openpyxl.load_workbook(path, data_only=False)
    assert workbook.sheetnames == [
        "Portada",
        "Datos",
        "Resumen",
        "Habilidades",
        "Empresas",
        "Calidad",
        "Diccionario",
        "Listas",
    ]
    assert workbook.active.title == "Portada"
    assert workbook["Listas"].sheet_state == "hidden"
    assert workbook.calculation.fullCalcOnLoad is True

    assert set(workbook["Datos"].tables) == {"tblDatos"}
    assert workbook["Datos"].tables["tblDatos"].ref == "A4:S16"
    assert set(workbook["Resumen"].tables) == {"tblPorCarrera", "tblPorEstado"}
    assert set(workbook["Habilidades"].tables) == {
        "tblHabilidades",
        "tblHabilidadesCarrera",
    }
    assert set(workbook["Empresas"].tables) == {"tblEmpresas", "tblSalariosCarrera"}
    assert set(workbook["Calidad"].tables) == {"tblCalidad"}
    assert set(workbook["Diccionario"].tables) == {"tblDiccionario"}

    assert workbook["Resumen"]["B6"].value == "=COUNTA(tblDatos[ID oferta])"
    assert workbook["Resumen"]["C9"].value.startswith("=COUNTIF")
    assert workbook["Calidad"]["B3"].value == "=AVERAGE(tblCalidad[Tasa de aprobación])"
    assert workbook["Habilidades"]["C5"].value.startswith("=B5 / COUNTA")
    assert workbook["Portada"]["A11"].value == "='Resumen'!B6"

    assert len(workbook["Resumen"].data_validations.dataValidation) == 1
    validation = workbook["Resumen"].data_validations.dataValidation[0]
    assert validation.formula1 == "=ListaCarreras"
    assert "C3" in str(validation.sqref)
    assert "ListaCarreras" in workbook.defined_names

    assert len(workbook["Resumen"]._charts) == 2
    assert len(workbook["Habilidades"]._charts) == 1
    assert len(workbook["Empresas"]._charts) == 2
    assert len(workbook["Calidad"]._charts) == 1
    assert len(workbook["Portada"]._charts) == 0

    assert len(workbook["Datos"].conditional_formatting) >= 2
    assert len(workbook["Calidad"].conditional_formatting) >= 2
    assert [str(item.sqref) for item in workbook["Portada"].conditional_formatting] == ["E11"]

    assert workbook["Calidad"].max_row >= 21
    assert workbook["Diccionario"].max_row >= 20
    assert workbook["Datos"]["B5"].value


def test_manifest_records_outputs(e2e_run: Path) -> None:
    manifest = json.loads((e2e_run / "run_manifest.json").read_text(encoding="utf-8"))
    assert manifest["status"] == "completed"
    assert manifest["quality"]["overall_score"] == 0.9427
    assert manifest["artifacts"]["dashboard_csv"] == "dashboard_dataset.csv"
    assert manifest["artifacts"]["excel"] == "practilizate.xlsx"
    assert manifest["outputs"]["dashboard_csv"].endswith("dashboard_dataset.csv")
    assert manifest["outputs"]["excel"].endswith("practilizate.xlsx")
    assert manifest["stages"]["export"]["csv_rows"] == 12
    assert manifest["stages"]["export"]["excel_sheets"] == 8


def test_export_command_rebuilds_deliverables(e2e_run: Path) -> None:
    from conftest import run_cli

    output_dir = e2e_run / "export-check"
    result = run_cli(
        "export",
        "--run-id",
        "e2e",
        "--artifacts-dir",
        str(e2e_run),
        "--output-dir",
        str(output_dir),
        "--format",
        "excel,csv",
    )
    assert result.returncode == 0, result.stderr
    summary = json.loads(result.stdout)
    assert summary["csv"]["rows"] == 12
    assert summary["excel"]["sheets"] == 8
    assert (output_dir / "dashboard_dataset.csv").is_file()
    assert (output_dir / "practilizate.xlsx").is_file()
